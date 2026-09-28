import re
import sys
from pathlib import Path

# Add python-docx installer check
try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:
    print("[INFO] python-docx not found. Installing it now...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def add_inline_formatting(paragraph, text):
    """Parses bold (**text**) and italic (*text*) formatting and adds to paragraph."""
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)

def convert_md_to_docx(md_path: Path, docx_path: Path):
    doc = docx.Document()
    
    # Page setup - Standard margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Arial'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    print(f"Reading Markdown from: {md_path}")
    md_content = md_path.read_text(encoding="utf-8")
    lines = md_content.split('\n')
    
    in_code_block = False
    code_text = []
    
    in_table = False
    table_rows = []
    
    in_list = False

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # 1. Code block handling
        if stripped.startswith('```') or stripped.startswith('%%'):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                
                # Format code run
                run = p.add_run("\n".join(code_text))
                run.font.name = 'Courier New'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                code_text = []
            else:
                in_code_block = True
            i += 1
            continue
            
        if in_code_block:
            code_text.append(line)
            i += 1
            continue
            
        # 2. Table handling
        if stripped.startswith('|'):
            in_table = True
            table_rows.append(stripped)
            i += 1
            continue
        elif in_table:
            # End of table, process it
            in_table = False
            # Clean separator lines
            clean_rows = []
            for r in table_rows:
                # Skip separator lines like |---|---|
                if re.match(r'^\|[\s:-|]*\|$', r):
                    continue
                cells = [c.strip() for c in r.split('|')[1:-1]]
                clean_rows.append(cells)
                
            if clean_rows:
                cols_count = len(clean_rows[0])
                table = doc.add_table(rows=len(clean_rows), cols=cols_count)
                table.style = 'Table Grid'
                
                # Populate table
                for r_idx, row_data in enumerate(clean_rows):
                    row = table.rows[r_idx]
                    for c_idx, val in enumerate(row_data):
                        # Ensure cell exists
                        if c_idx < len(row.cells):
                            cell = row.cells[c_idx]
                            cell.text = "" # Clear default
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(4)
                            p.paragraph_format.space_after = Pt(4)
                            p.paragraph_format.left_indent = Inches(0.05)
                            
                            # Bold headers
                            if r_idx == 0:
                                run = p.add_run(val)
                                run.bold = True
                                run.font.size = Pt(10.5)
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                set_cell_background(cell, "336699") # Dark Blue Header
                            else:
                                add_inline_formatting(p, val)
                                run = p.runs[0] if p.runs else p.add_run("")
                                run.font.size = Pt(10)
                                if r_idx % 2 == 0:
                                    set_cell_background(cell, "F2F5F9") # Zebra striping
                                    
                doc.add_paragraph() # Spacer
            table_rows = []
            # Do not increment i, let current line be processed

        # 3. Headings
        if stripped.startswith('# '):
            heading = doc.add_heading(level=0)
            heading.paragraph_format.space_before = Pt(12)
            heading.paragraph_format.space_after = Pt(6)
            run = heading.add_run(stripped[2:])
            run.font.name = 'Arial'
            run.font.color.rgb = RGBColor(0x11, 0x44, 0x77)
            run.bold = True
        elif stripped.startswith('## '):
            heading = doc.add_heading(level=1)
            heading.paragraph_format.space_before = Pt(18)
            heading.paragraph_format.space_after = Pt(6)
            run = heading.add_run(stripped[3:])
            run.font.name = 'Arial'
            run.font.color.rgb = RGBColor(0x22, 0x55, 0x88)
            run.bold = True
        elif stripped.startswith('### '):
            heading = doc.add_heading(level=2)
            heading.paragraph_format.space_before = Pt(12)
            heading.paragraph_format.space_after = Pt(4)
            run = heading.add_run(stripped[4:])
            run.font.name = 'Arial'
            run.font.color.rgb = RGBColor(0x33, 0x66, 0x99)
            run.bold = True
        elif stripped.startswith('#### '):
            heading = doc.add_heading(level=3)
            heading.paragraph_format.space_before = Pt(6)
            heading.paragraph_format.space_after = Pt(2)
            run = heading.add_run(stripped[5:])
            run.font.name = 'Arial'
            run.font.color.rgb = RGBColor(0x44, 0x77, 0xAA)
            run.bold = True
            
        # 4. Bullet lists
        elif stripped.startswith('* ') or stripped.startswith('- ') or (stripped.startswith('1. ') and not re.match(r'^\d+\.\d+', stripped)):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            # Remove prefix
            prefix_len = 2 if stripped.startswith('* ') or stripped.startswith('- ') else 3
            add_inline_formatting(p, stripped[prefix_len:])
            
        # 5. Regular paragraph or empty line
        else:
            if stripped == "" or stripped == "---":
                if stripped == "---":
                    doc.add_paragraph().paragraph_format.space_after = Pt(12)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                add_inline_formatting(p, stripped)
                
        i += 1

    doc.save(str(docx_path))
    print(f"[SUCCESS] Saved Word Document to: {docx_path}")

if __name__ == "__main__":
    root_dir = Path(__file__).resolve().parents[1]
    md_file = root_dir / "Raporti_HealthConnect_AI.md"
    docx_file = root_dir / "Raporti_HealthConnect_AI.docx"
    convert_md_to_docx(md_file, docx_file)
